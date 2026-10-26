"""Deterministic product rendering and quality checks."""
from __future__ import annotations

from pathlib import Path
from typing import Any

from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, PageBreak


def render_product_files(output_dir: Path, title: str, goal: str, blueprint: dict[str, Any]) -> dict[str, Any]:
    output_dir.mkdir(parents=True, exist_ok=True)
    safe_title = "".join(char if char.isalnum() or char in "-_" else "_" for char in title).strip("_") or "pinrin_product"
    pdf_path = output_dir / f"{safe_title}.pdf"
    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"]), Spacer(1, .2 * inch), Paragraph(goal, styles["BodyText"]), Spacer(1, .25 * inch)]
    outline = blueprint.get("outline", []) if isinstance(blueprint, dict) else []
    if not isinstance(outline, list):
        outline = []
    if not outline:
        outline = ["Introduction", "Core framework", "Exercises and checklist", "Next steps"]
    for index, section in enumerate(outline, 1):
        if isinstance(section, dict):
            heading = section.get("title") or section.get("name") or f"Section {index}"
            body = section.get("description") or section.get("content") or "Section content to review."
        else:
            heading, body = str(section), "Section content to review."
        story.extend([Paragraph(f"{index}. {heading}", styles["Heading2"]), Paragraph(str(body), styles["BodyText"]), Spacer(1, .18 * inch)])
    SimpleDocTemplate(str(pdf_path), pagesize=LETTER, rightMargin=.7 * inch, leftMargin=.7 * inch, topMargin=.7 * inch, bottomMargin=.7 * inch).build(story)
    files = [{"path": str(pdf_path), "kind": "generated_pdf"}]
    try:
        from docx import Document
        document = Document()
        document.add_heading(title, 0)
        document.add_paragraph(goal)
        for index, section in enumerate(outline, 1):
            heading = section.get("title") if isinstance(section, dict) else str(section)
            document.add_heading(f"{index}. {heading}", level=1)
            document.add_paragraph(str(section.get("content", "Section content to review.")) if isinstance(section, dict) else "Section content to review.")
        docx_path = output_dir / f"{safe_title}.docx"
        document.save(docx_path)
        files.append({"path": str(docx_path), "kind": "generated_docx"})
    except ImportError:
        pass
    return {"status": "generated", "title": title, "files": files, "outline": outline}


def validate_product_files(files: list[dict[str, Any]], images: list[str] | None = None) -> dict[str, Any]:
    checks: list[dict[str, Any]] = []
    for item in files:
        path = Path(item["path"])
        exists = path.exists() and path.stat().st_size > 0
        check = {"file": str(path), "exists": exists, "bytes": path.stat().st_size if exists else 0}
        if path.suffix.lower() == ".pdf" and exists:
            try:
                from pypdf import PdfReader
                reader = PdfReader(str(path))
                check["pages"] = len(reader.pages)
                check["text_present"] = any((page.extract_text() or "").strip() for page in reader.pages)
                check["valid"] = bool(reader.pages) and check["text_present"]
            except Exception as exc:
                check.update({"valid": False, "error": str(exc)})
        else:
            check["valid"] = exists
        checks.append(check)
    image_checks = []
    for image_name in images or []:
        path = Path(image_name)
        item = {"file": str(path), "exists": path.exists()}
        if path.exists():
            try:
                from PIL import Image
                with Image.open(path) as image:
                    item.update({"width": image.width, "height": image.height, "valid": image.width > 0 and image.height > 0})
            except Exception as exc:
                item.update({"valid": False, "error": str(exc)})
        else:
            item["valid"] = False
        image_checks.append(item)
    all_valid = all(item.get("valid", False) for item in checks + image_checks) if checks + image_checks else False
    return {"status": "passed" if all_valid else "failed", "files": checks, "images": image_checks}
